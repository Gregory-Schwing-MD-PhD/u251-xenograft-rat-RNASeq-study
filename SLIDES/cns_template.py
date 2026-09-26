# -*- coding: utf-8 -*-
"""Slide plumbing for the 2026 CNS speaker template, vendored into this repository so the Abstract 418 deck builds
from the u251 directory alone. Copied (unchanged in behaviour) from the SchwingNet CNS deck generators:
make_cns_deck_v2.fit / img_size, make_cns_deck_v4 (geometry, palette, add_text, picture, content_slide, refs, arrow,
navy_bar) and make_cns_deck_v6.notes.

The template itself is the CNS's file and is not committed (this repository is public): put
'2026 CNS AM Speaker Powerpoint Template.pptx' in SLIDES/templates/ (gitignored). Template geometry (16:9): layout 0
"Title Slide"; layout 2 "1_Title and Content Vertical" (navy band at the left, white content area). A text box with
zero or negative extent makes PowerPoint refuse the whole file, so add_text clamps extents.
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

TEMPLATE = Path(__file__).resolve().parent / "templates" / "2026 CNS AM Speaker Powerpoint Template.pptx"

NAVY = RGBColor(0x19, 0x1D, 0x63)
BLUE = RGBColor(0x4C, 0x70, 0xB7)
GOLD = RGBColor(0xFF, 0xC4, 0x2F)
RED = RGBColor(0xEF, 0x46, 0x27)
GREY = RGBColor(0x7A, 0x7A, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

L = 1750000
R = 11900000
TOP = 1420000
BOTTOM = 6250000
REFS_TOP = 6380000
GAP = 300000
BODY = 18
CAP = 12
REF = 10


def img_size(path):
    with Image.open(path) as im:
        return im.size


def fit(path, max_w, max_h):
    """Width and height (EMU) that fit the image inside max_w x max_h, aspect kept."""
    w, h = img_size(path)
    scale = min(max_w / w, max_h / h)
    return int(w * scale), int(h * scale)


def add_text(slide, left, top, width, height, lines, size=BODY, color=NAVY, bold=False, gap=8, align=None, anchor=None):
    height = max(int(height), 250000)
    width = max(int(width), 250000)
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    if anchor:
        tf.vertical_anchor = anchor
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(line, list):          # coloured runs: [(text, colour), ...]
            for text, col in line:
                r = p.add_run()
                r.text = text
                r.font.size = Pt(size)
                r.font.color.rgb = col
                r.font.bold = bold
        else:
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
        p.space_before = Pt(0)
        p.space_after = Pt(gap)
        if align:
            p.alignment = align
    return tb


def picture(slide, path, left, top, width, height):
    return slide.shapes.add_picture(str(path), Emu(left), Emu(top), width=Emu(width), height=Emu(height))


def content_slide(prs, title, size=30):
    s = prs.slides.add_slide(prs.slide_layouts[2])
    for ph in list(s.placeholders):
        if ph.placeholder_format.idx != 0:
            ph._element.getparent().remove(ph._element)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Emu(L), Emu(300000), Emu(R - L), Emu(1000000)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM
    tf.text = title
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
            r.font.bold = True
            r.font.color.rgb = NAVY
    return s


def refs(slide, text):
    add_text(slide, L, REFS_TOP, R - L, 420000, [text], size=REF, color=GREY, gap=0)


def arrow(slide, left, top):
    a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(left), Emu(top), Emu(320000), Emu(320000))
    a.fill.solid()
    a.fill.fore_color.rgb = GOLD
    a.line.color.rgb = GOLD
    a.shadow.inherit = False


def navy_bar(slide, left, top, width, height, text, size=18):
    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.color.rgb = NAVY
    bar.shadow.inherit = False
    tf = bar.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size if i == 0 else size - 4)
        r.font.bold = i == 0
        r.font.color.rgb = WHITE
    return bar


def notes(slide, lines):
    """Speaker notes: the text taken off the slide."""
    slide.notes_slide.notes_text_frame.text = "\n\n".join(lines)
