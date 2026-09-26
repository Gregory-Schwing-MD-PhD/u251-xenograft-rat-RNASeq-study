# -*- coding: utf-8 -*-
"""The CNS 2026 Abstract 418 talk (Characterizing the Transcriptomic Recurrence Signature of Glioblastoma Following
Laser Interstitial Thermal Therapy; presentation SSTU02, Sunrise Science Session, Tumor 2, Monday 2 November 2026,
7:00-7:06 AM, Room 147B) on the official 2026 CNS speaker template, in the style of the SchwingNet deck: one large
figure per slide, the result and its explanation in a few lines, a citation line with the study design on every data
slide, no logos, no company names (ACCME guidelines for CNS speakers; the disclosures are read aloud from the title
slide's notes, CNS displays the list).

Version 3 (2026-09-26), after an adversarial check of every claim on version 2 against the manuscript, the
supplementary sheets, the pipeline outputs and the cited papers (104 findings confirmed). What changed:
  - every statement is nominal or FDR as it really is; only translation initiation clears FDR in the screen;
  - the ablation is monitored by diffusion-weighted MRI (MR thermometry could not run on the animal scanner);
  - human AND shared reads enter the analysis; one rat-brain control holds tumour cells;
  - the contamination slide argues magnitude, not direction (07_contamination_magnitude.py);
  - respiration does not agree across gene-set collections, so 'energy' became glycolysis, and cholesterol synthesis
    rising is shown; the theme chart has its own slide and the translation slide leads with the running-sum plot;
  - the subtype q values are a correct Benjamini-Hochberg (the supplement's column is mis-ordered);
  - the funnel stops at the 15 compounds that clear both barrier models; the prior-art audit is the top 20 of 54 and
    includes thioridazine; the candidate table carries overall ranks;
  - the two strongest reversals are dimethyloxalylglycine and deferoxamine;
  - references are numbered by first appearance (Refs below), Nagaraja 2026 is J Neurosurg 145:364-377.

Every number on a slide is read from figures/facts.json, the figures_cns/*.json and *.csv written by the figure
scripts, or the DepMap table. Figures: 04_make_cns_figures.py, 05_make_more_figures.py, 06_contamination_check.py,
07_contamination_magnitude.py.

    python SLIDES/04_build_cns_template_deck.py [--out PATH]    (default SLIDES/CNS2026_Schwing_Abstract418_CNStemplate.pptx)
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))
from u251_paths import GSEA_LOO  # noqa: E402
from cns_template import (BLUE, BOTTOM, CAP, GAP, GREY, L, NAVY, R, REF, REFS_TOP, TEMPLATE, TOP, WHITE,  # noqa: E402
                          add_text, arrow, content_slide, fit, navy_bar, notes, picture, refs)

FIG = HERE / "figures_cns"
DECK = {r["n"]: r for r in json.load(open(HERE / "deck_2026-09-08_text.json", encoding="utf-8"))}
F = json.load(open(HERE / "figures" / "facts.json", encoding="utf-8"))
TH = F["themes"]
J = lambda name: json.load(open(FIG / name, encoding="utf-8"))  # noqa: E731
TITLE = "Characterizing the Transcriptomic Recurrence Signature of Glioblastoma Following Laser Interstitial Thermal Therapy"
NB = lambda s: s.replace(" ", " ")  # noqa: E731   a name never breaks across lines
AUTHORS = ["  ·  ".join([NB("Gregory J. Schwing, MD, PhD¹"), NB("Tavarekere N. Nagaraja, PhD²")]),
           "  ·  ".join([NB("Indrani Datta, DHI²"), NB("Ian Y. Lee, MD²")])]
AFFIL = ["¹ Department of General Surgery, Detroit Medical Center and Wayne State University, Detroit, MI",
         "² Department of Neurosurgery, Hermelin Brain Tumor Center, Henry Ford Health, Detroit, MI"]
SESSION = ["Presentation SSTU02 (Abstract 418)  ·  Sunrise Science Session, Tumor 2",
           "Monday 2 November 2026, 7:00–7:06 AM",
           "Room 147B, Walter E. Washington Convention Center"]
DESIGN = ("Orthotopic U251N glioblastoma xenograft in the athymic rat, MRI-guided LITT; bulk RNA-seq of primary and recurrent "
          "tumors, n = 3 per arm; GEO GSE338105.")

# ---------------------------------------------------------------------- references, numbered by first appearance
REFS = {
    "chen": "Chen C, et al. J Neurooncol 2021;151:429–442.",
    "cleary": "Cleary RT, et al. Neuro Oncol 2026;28:1649–1661.",
    "tao": "Tao R, et al. J Immunol 2026;215:vkaf327.",
    "nag26": "Nagaraja TN, et al. J Neurosurg 2026;145:364–377.",
    "nag21": "Nagaraja TN, et al. Acta Neurochir 2021;163:3455–3463.",
    "xengsort": "Zentgraf J, Rahmann S. Algorithms Mol Biol 2021;16:2.",
    "maleki": "Maleki F, et al. Hum Genomics 2019;13(Suppl 1):42.",
    "neftel": "Neftel C, et al. Cell 2019;178:835–849.",
    "garofano": "Garofano L, et al. Nat Cancer 2021;2:141–156.",
    "yoo": "Yoo M, et al. Bioinformatics 2015;31:3069–3071.",
    "risso": "Risso D, et al. Nat Biotechnol 2014;32:896–902.",
    "venkatesh": "Venkatesh HS, et al. Nature 2019;573:539–545.",
    "venkataramani": "Venkataramani V, et al. Nature 2019;573:532–538.",
    "swanson": "Swanson K, et al. Bioinformatics 2024;40:btae416.",
    "daina": "Daina A, Zoete V. ChemMedChem 2016;11:1117–1121.",
    "minden": "Minden MD, et al. Am J Hematol 2014;89:363–368.",
    "su": "Su Z, et al. Cell Death Dis 2021;12:251.",
    "boursi": "Boursi B, et al. Pharmacoepidemiol Drug Saf 2016;25:1179–1185.",
    "drljaca": "Drljača J, et al. CNS Neurosci Ther 2022;28:1447–1457.",
    "depmap": "DepMap 24Q4 (depmap.org), Chronos gene effect, U-251 MG (ACH-000232); Dempster JM, et al. Genome Biol 2021;22:343.",
    "sun": "Sun S, et al. J Transl Med 2025;23:25.",
}
SUP = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")


class Refs:
    """cite(*keys) gives the superscript and numbers a reference the first time it is cited; line() lists the
    references cited on the current slide, in number order, and starts the next slide."""

    def __init__(self):
        self.num, self.page = {}, []

    def __call__(self, *keys):
        out = []
        for k in keys:
            assert k in REFS, k
            if k not in self.num:
                self.num[k] = len(self.num) + 1
            if k not in self.page:
                self.page.append(k)
            out.append(str(self.num[k]).translate(SUP))
        return "˒".join(out)

    def only(self, *keys):             # cited on the slide without a superscript (author-year inside a figure)
        self(*keys)

    def line(self, tail=""):
        keys = sorted(self.page, key=lambda k: self.num[k])
        self.page = []
        body = "   ".join(f"{self.num[k]}.\u00a0{REFS[k]}" for k in keys)
        return (body + ("   " if body and tail else "") + tail).strip()


R_ = Refs()


def old_notes(*ns):
    return [DECK[n]["notes"].strip() for n in ns if DECK[n]["notes"].strip()]


def m(v, signed=False, nd=2):
    """A number with a typographic minus, so no line breaks after a hyphen."""
    s = f"{v:+.{nd}f}" if signed else f"{v:.{nd}f}"
    return s.replace("-", "−")


def pv(p):
    return f"{p:.3f}" if p < 0.01 else f"{p:.2f}"


def head_box(slide, left, top, width, height, head, sub, size=20):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
    box.fill.solid(); box.fill.fore_color.rgb = NAVY; box.line.color.rgb = NAVY; box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = head; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = WHITE
    if sub:
        add_text(slide, left, top + height + 80000, width, 800000, sub, size=15, color=NAVY, gap=0, align=PP_ALIGN.CENTER)


def fig_and_text(s, img, lines, img_w, size=16, gap=10, caption=None, text_top=0):
    """A figure at the left up to img_w wide and the column height, the explanation at the right."""
    w, h = fit(img, img_w, BOTTOM - TOP - (450000 if caption else 0))
    picture(s, img, L, TOP, w, h)
    if caption:
        add_text(s, L, TOP + h + 60000, w, 420000, [caption], size=CAP, color=GREY, gap=0)
    x2 = L + w + GAP
    add_text(s, x2, TOP + text_top, R - x2, BOTTOM - TOP - text_top, lines, size=size, gap=gap)
    return w, h


def fig_top(s, img, lines, reserve, size=13, gap=5, color=NAVY):
    """A figure centred across the top, the explanation under it."""
    w, h = fit(img, R - L, BOTTOM - TOP - reserve)
    picture(s, img, L + (R - L - w) // 2, TOP, w, h)
    add_text(s, L, TOP + h + 60000, R - L, BOTTOM - (TOP + h + 60000), lines, size=size, color=color, gap=gap)
    return w, h


def flat(p):
    """No inherited indent on a placeholder paragraph."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0"); pPr.set("indent", "0")


def timed(t, *rest):
    return [t] + [x for x in rest if x]


def build(out: Path):
    C = R_
    prs = Presentation(str(TEMPLATE))
    sld = prs.slides._sldIdLst
    RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    for e in list(sld):
        sld.remove(e)
        prs.part.drop_rel(e.get(RID))

    # data read once
    srt = list(csv.DictReader(open(FIG / "chart_sorting.csv", encoding="utf-8")))
    grp = lambda coh, col="graft": [float(r[col]) for r in srt if r["cohort"].startswith(coh)]  # noqa: E731
    pri, rec, ctl, cul = grp("Primary"), grp("Recurrent"), grp("Control"), grp("In vitro")
    tum_both = grp("Primary", "both_pct") + grp("Recurrent", "both_pct")
    tum_assigned = [float(r["graft"]) + float(r["host"]) for r in srt if r["cohort"].startswith(("Primary", "Recurrent"))]
    h_pri, h_rec = sum(pri) / 3, sum(rec) / 3
    rs = J("chart_running_sum.json")
    cc = J("contamination_check.json")
    mag = J("contamination_magnitude.json")
    # leave-one-tumour-out GSEA (ANALYSIS/gsea_leave_one_out, 41 runs of the published command; the full seed-1234 run
    # reproduces the published report exactly)
    assert json.load(open(GSEA_LOO / "reproduction.json"))["exact"] is True
    loo = list(csv.DictReader(open(GSEA_LOO / "loo_sets.tsv", encoding="utf-8"), delimiter="\t"))
    loo_scr = list(csv.DictReader(open(GSEA_LOO / "loo_screen.tsv", encoding="utf-8"), delimiter="\t"))
    assert len({(r["condition"], r["seed"]) for r in loo}) == 41
    TIK = "KEGG_MEDICUS_REFERENCE_TRANSLATION_INITIATION"
    ti_loo = [r for r in loo if r["set"] == TIK and not r["condition"].endswith("_nofilter")]
    qs = lambda cond: [float(r["fdr_q"]) for r in ti_loo if r["condition"] == cond]  # noqa: E731
    q_full, q_68 = qs("full"), qs("drop_IL68B")
    assert len(q_full) == len(q_68) == 5
    never = [c.replace("drop_", "") for c in dict.fromkeys(r["condition"] for r in ti_loo) if c != "full" and min(qs(c)) >= 0.05]
    ti_all = [r for r in loo if r["set"] == TIK]
    ti_pmax = max(float(r["nom_p"]) for r in ti_all)
    ti_nes = [float(r["NES"]) for r in ti_all]
    lead_pmax = max(float(r["nom_p"]) for r in loo if r["lead"] == "True")
    n_full_ok = sum(q < 0.05 for q in q_full)
    full_up25 = max(int(r["up_q25"]) for r in loo_scr if r["condition"] == "full")
    assert ti_pmax < 0.001 and full_up25 == 0
    sets_p = J("chart_gsea_sets.json")
    thr = J("chart_threshold_sweep.json")
    themes = {r["set"]: r for r in J("chart_themes.json")}
    sub = J("chart_subtypes.json")
    ppi = J("chart_ppi.json")
    fun = J("chart_funnel.json")
    ds = J("chart_drug_scatter.json")
    s13 = J("s13_top20.json")
    tiers = J("table_prior_art.json")
    tj = J("chart_trajectory_pca.json")
    dm = {r["gene"]: r for r in csv.DictReader(open(ROOT / "depmap" / "depmap_u251_targets.csv", encoding="utf-8"))}
    K = {"init": "KEGG_MEDICUS_REFERENCE_TRANSLATION_INITIATION", "elong": "REACTOME_EUKARYOTIC_TRANSLATION_ELONGATION", "rib": "KEGG_RIBOSOME",
         "gcn2": "REACTOME_RESPONSE_OF_EIF2AK4_GCN2_TO_AMINO_ACID_DEFICIENCY", "starv": "REACTOME_CELLULAR_RESPONSE_TO_STARVATION"}
    top = F["gsea_down_top"]

    # ---- 1 title
    s = prs.slides.add_slide(prs.slide_layouts[0])
    ph = {p.placeholder_format.idx: p for p in s.placeholders}
    t = ph[0]
    t.left, t.top, t.width, t.height = Emu(5185558), Emu(950000), Emu(6600000), Emu(2000000)
    t.text_frame.word_wrap = True
    t.text_frame.text = TITLE
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = WHITE
    sb = ph[1]
    sb.left, sb.top, sb.width, sb.height = Emu(5185558), Emu(3050000), Emu(6600000), Emu(1350000)
    sb.text_frame.word_wrap = True
    sb.text_frame.text = AUTHORS[0]
    for a in AUTHORS[1:] + AFFIL:
        p = sb.text_frame.add_paragraph(); p.text = a
    for i, p in enumerate(sb.text_frame.paragraphs):
        p.space_after = Pt(1 if i == 0 else 5 if i == 1 else 1)
        p.level = 0; flat(p)
        for r in p.runs:
            r.font.size = Pt(15 if i < 2 else 11); r.font.color.rgb = WHITE
    body = ph[10]
    body.left, body.top, body.width, body.height = Emu(5185558), Emu(4550000), Emu(6600000), Emu(1450000)
    body.text_frame.word_wrap = True
    body.text_frame.text = "CNS 2026 Annual Meeting, Washington, DC"
    for line in SESSION:
        p = body.text_frame.add_paragraph(); p.text = line
    for p in body.text_frame.paragraphs:
        p.space_after = Pt(3)
        p.level = 0; flat(p)
        for r in p.runs:
            r.font.size = Pt(12); r.font.color.rgb = WHITE
    notes(s, [
        "0:00–0:08  Title. Name, institutions, and that this is a rat model of MRI-guided LITT. Do not read the title aloud.",
        "Disclosures, read aloud (CNS displays the list; no disclosure slide per the ACCME speaker guidelines): I.Y.L. has consulting "
        "agreements with Medtronic, Inc. (Minneapolis, MN) and Monteris Medical, Inc. (Plymouth, MN). All other authors declare no "
        "conflicts of interest. Supported by a Henry Ford Health Physician Scientist Award A20050 (I.Y.L.). No commercial support was "
        "received for this analysis. Drug candidates discussed here are computational predictions. None is approved for glioblastoma "
        "and none has been tested in this model.",
        "Why the disclosures matter here: Medtronic makes Visualase, the LITT system used in this rat model, and Monteris makes "
        "NeuroBlate, a competing system, so the room will want to hear both named. Some agents that appear later (ifosfamide, "
        "ganciclovir, progesterone, metformin) have been studied in glioblastoma; ciclopirox has not been tested in patients with "
        "glioblastoma.",
        "Timing: 25 slides in six minutes, about 14 s each. If running long, cut in this order: 8 (dish to brain), 21 (prior-art "
        "table), 19 (scatter); then 17 (hubs; also drop the network from beat 4 of the roadmap) and 14 (what else moves; also drop "
        "'(next slide)' on slide 13 and the hypoxia beat of the conclusions). Each saves 8 to 12 s. Keep 15 and 24: the one-tumor "
        "caveat, the methylation result and the next steps are said there."])
    C.line()

    # ---- 2 where this is going: the five beats and the study in one strip
    s = content_slide(prs, "Where this is going")
    items = ["The problem: LITT kills the core; the cells that come back sat in the sublethal margin",
             "The model: MRI-guided LITT in an orthotopic rat xenograft, primary and recurrent tumors, three per arm",
             "The finding: the recurrent tumor turns its translation machinery down; hypoxia signalling falls too, nominally",
             "The context: published subtype scores, and the network the changed genes form",
             "The opening: which existing drug, predicted to cross the blood–brain barrier, opposes that state, and whether the tumor depends on its target"]
    y = TOP
    for i, txt in enumerate(items, 1):
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(L), Emu(y), Emu(400000), Emu(400000))
        box.fill.solid(); box.fill.fore_color.rgb = NAVY; box.line.color.rgb = NAVY; box.shadow.inherit = False
        tf = box.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i); r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
        add_text(s, L + 540000, y - 30000, R - L - 540000, 500000, [txt], size=15, gap=0)
        y += 520000
    img = FIG / "overview_pipeline.png"
    w, h = fit(img, R - L, BOTTOM - y - 420000)
    picture(s, img, L + (R - L - w) // 2, y + 100000, w, h)
    add_text(s, L, y + 100000 + h + 40000, R - L, 300000,
             ["Then: differential expression, gene-set enrichment, published subtype signatures, a STRING network, and the drug gene sets that oppose the recurrence signature."],
             size=CAP, color=GREY, gap=0)
    notes(s, ["0:08–0:20  Roadmap. Five beats, twelve seconds; do not linger.",
              "The strip is the study in one line: the human line, the orthotopic tumor, the ablation (the frame is from the video three slides on), the recurrence, and the sequencing of three primary and three recurrent tumors."])
    C.line()

    # ---- 3 LITT ablates the core; recurrence starts at the margin
    s = content_slide(prs, "LITT ablates the core; recurrence starts at the margin")
    fig_and_text(s, FIG / "fig_margin.png", [
        f"LITT is increasingly used for deep-seated and recurrent glioblastoma, where open resection carries high morbidity.{C('chen')}",
        "Thermal dose falls off with distance from the fiber. Beyond the ablative threshold lies a sublethal margin.",
        "The cells that seed recurrence are the ones that survive there.",
        "So the question is not whether ablation works at the core. It is what the cells that survived the margin look like when they regrow."],
        img_w=6400000, size=16, gap=12, text_top=100000)
    refs(s, C.line("Diagram: thermal dose falls with distance from the fiber; schematic, not to scale."))
    notes(s, ["0:20–0:38  Background. The one idea to land: LITT kills the core, but the cells that come back sat in the sublethal margin. Point at the diagram once."])

    # ---- 4 what that margin is known to do, and what has not been done
    s = content_slide(prs, "What is known about the margin, and what is not")
    img = FIG / "fig_margin_known.png"
    w, h = fit(img, R - L, BOTTOM - TOP - 900000)
    picture(s, img, L + (R - L - w) // 2, TOP, w, h)
    C.only("cleary", "tao", "nag26")
    navy_bar(s, L, TOP + h + 120000, R - L, 760000,
             "Does post-LITT recurrence carry its own transcriptional program, and does that program point to a drug?", size=18)
    refs(s, C.line())
    notes(s, ["0:38–0:55  The gap. The vessels and the immune side of that margin have been described, and a first-pass comparison of recurrent against primary on these same libraries found cell-cycle, motility and inflammatory genes up in recurrence (the J Neurosurg abstract reports four per arm; our Methods say four primary and three recurrent, the fourth being the failed graft now used as a control: reconcile before the talk). What has not been done is a read of the tumor's own program, sorted from host reads and read gene-set-wide, and whether it points to a drug. End on the question and pause."])

    # ---- 5 the model, with the ablation video
    s = content_slide(prs, "The model: MRI-guided LITT in a rat xenograft")
    gapc = 360000
    colw = (R - L - 2 * gapc) // 3
    xs = [L, L + colw + gapc, L + 2 * (colw + gapc)]
    heads = ("U251N xenograft", "MRI-guided LITT", "Tumors harvested")
    subs = (["orthotopic human glioblastoma line in the athymic RNU/RNU rat"],
            ["ablation of the primary tumor, watched by diffusion-weighted MRI"],
            ["primary before ablation, recurrent after regrowth, three animals per arm"])
    for x, head, sb_ in zip(xs, heads, subs):
        head_box(s, x, TOP, colw, 480000, head, sb_, size=20)
    for x in xs[:2]:
        arrow(s, x + colw + (gapc - 320000) // 2, TOP + 80000)
    vy = TOP + 1300000
    vs = BOTTOM - vy
    s.shapes.add_movie(str(FIG / "ablation_512.mp4"), Emu(L), Emu(vy), Emu(vs), Emu(vs), poster_frame_image=str(FIG / "video_poster.png"), mime_type="video/mp4")
    x2 = L + vs + GAP
    add_text(s, x2, vy, R - x2, 2700000, [
        f"A model the Henry Ford group established and characterized.{C('nag26', 'nag21')} The ablation is watched in real time by diffusion-weighted MRI and the laser is stopped when the diffusion lesion covers the tumor; MR thermometry could not run on the animal scanner.",
        f"The graft is human and the host is rat, which is the point: about nine reads in ten ({min(tum_assigned):.0f} to {max(tum_assigned):.0f} %) can be assigned to one species, so the tumor compartment can be read largely on its own.",
        "Primary tumors were harvested before ablation and recurrent tumors after regrowth; the primary-versus-recurrent comparison rests on those six transcriptomes."],
        size=14.5, gap=8)
    add_text(s, x2, vy + 2750000, R - x2, 650000,
             ["Video: an ablation under diffusion-weighted MRI in this model; click to play (7.9 s)."],
             size=CAP, color=GREY, gap=0)
    refs(s, C.line())
    notes(s, ["0:55–1:12  The model. Move fast. The one thing to stress is that the graft is human and the host is rat, which is what makes the tumor compartment separable. The laser (Visualase) is stopped when the diffusion-weighted lesion covers the tumor.",
              "The video is the intra-MRI LITT with diffusion imaging from the 2021 technical note (Ablation video_TNN_18Sept26.pptx); the source clip is 64 x 64 px, re-encoded at 512 px for the slide."])

    # ---- 6 ten libraries, sorted by species
    s = content_slide(prs, "Ten libraries: what fraction of each is human")
    fig_and_text(s, FIG / "chart_sorting.png", [
        f"Every read is classified by its sequence as human, rat, shared by both genomes, or unresolved before anything is counted.{C('xengsort')}",
        f"The culture sample is {cul[0]:.0f} % human. In the brain the tumors are {min(pri):.0f} to {max(pri):.0f} % human before ablation and {min(rec):.0f} to {max(rec):.0f} % after. Human reads enter the analysis together with the reads shared by both genomes, a further {min(tum_both):.0f} to {max(tum_both):.0f} % of each tumor library; the shared reads are the route by which rat signal could leak in.",
        f"The three rat-brain controls went through the identical pipeline. Their human reads ({min(ctl):.1f} to {max(ctl):.1f} %) show what the human stream picks up in brain without an established tumor, and anchor the contamination check.",
        [(f"The recurrent tumors carry more rat than the primaries ({h_rec:.0f} % against {h_pri:.0f} % human): a confound to test, not to ignore.", BLUE)]],
        img_w=6500000, size=13.5, gap=8, text_top=40000)
    refs(s, C.line("k = 25; percentages from the classification logs. " + DESIGN))
    notes(s, ["1:12–1:25  About 13 s. The point is the graft-fraction gradient: primary about half human, recurrent about a third. Say plainly that this is a confound and that it is tested on the contamination slide. The culture sample and the controls are not in the differential model.",
              "If asked about the controls: one of them, N269B, carries human Y-linked transcripts of the male U251 line, so it holds some tumor cells or carryover; it is a contamination anchor, not a pure rat brain."])

    # ---- 7 the analysis, explained
    s = content_slide(prs, "The analysis, and how to read the numbers that follow")
    heads = ("Sort", "Test", "Interpret")
    subs = (["reads sorted into human, rat, shared and unresolved (xengsort, k = 25)"],
            ["DESeq2 Wald test with ashr shrinkage; FDR < 0.05, |log₂ fold change| > 1"],
            ["GSEA enrichment, GSVA subtypes, STRING network, DSigDB drug gene sets"])
    for x, head, sb_ in zip(xs, heads, subs):
        head_box(s, x, TOP, colw, 480000, head, sb_, size=20)
    for x in xs[:2]:
        arrow(s, x + colw + (gapc - 320000) // 2, TOP + 80000)
    add_text(s, L, TOP + 1450000, R - L, BOTTOM - TOP - 1450000, [
        "Enrichment asks whether a whole gene set drifts toward the down- or up-regulated end of the ranked transcriptome. The normalized enrichment score (NES) is that drift; negative means down in recurrence.",
        f"Significance is permuted on gene-set membership, not on samples: with three per arm there are too few sample permutations. Gene-set permutation is more reproducible than sample permutation at this size{C('maleki')}, but it is anti-conservative for correlated sets, so nominal p is shown beside FDR q throughout.",
        f"Subtypes are scored per tumor by GSVA over the published Neftel{C('neftel')} and Garofano{C('garofano')} signatures. Drugs are ranked by how strongly each compound's DSigDB gene set{C('yoo')} sits among the genes that fall in recurrence, then weighted by predicted blood–brain barrier permeability.",
        "Raw sequencing data are public: GEO accession GSE338105."], size=14.5, gap=8)
    refs(s, C.line())
    notes(s, ["1:25–1:40  The analysis. Enrichment was permuted on gene-set membership, not on samples, because n = 3. If someone challenges the statistics later, this slide is the answer. Mention the GEO accession is already public.",
              "If asked what a DSigDB drug set is: for the leading hits, including ciclopirox, it is the drug's expression signature, the genes it induces in cultured cells; those genes are lower in recurrence."])

    # ---- 8 from the dish to the brain, then after ablation (exploratory)
    s = content_slide(prs, "From the dish to the brain, then after ablation")
    fig_and_text(s, FIG / "chart_trajectory_pca.png", [
        "Exploratory: the culture sample beside the six tumors, log₂ counts per million over the 500 most variable genes. The culture sample is one library and sits outside every test in this talk.",
        f"Culture separates from the tumors on the first axis ({tj['pc1']:.1f} % of the variance in those genes); primary and recurrent separate mainly on the second ({tj['pc2']:.1f} %).",
        f"Across those genes the culture sample is {tj['dist_culture_to_tumour_centroid']:.0f} log₂ units from the tumor centroid, about {tj['ratio']:.0f} times the {tj['dist_primary_to_recurrent_centroid']:.0f} between the primary and recurrent centroids: growth in the brain moves the transcriptome more than ablation and regrowth do.",
        [("Part of the culture distance may be rat signal: the culture library carries none, the tumor libraries carry some.", BLUE)]],
        img_w=5600000, size=14.5, gap=10, text_top=80000)
    refs(s, C.line("Principal components of log2 CPM, genes with mean CPM ≥ 1, 500 most variable, seven libraries; Euclidean distances between centroids in the same 500 genes; exploratory, outside the differential model. " + DESIGN))
    notes(s, ["1:40–1:48  CUT FIRST if running long. The repository's framing (Camphausen 2005: the brain microenvironment dominates the expression of glioma lines) on our own data. One culture library, so no statistics; the manuscript's PCA is the six tumors alone on the next slide. On the first axis the recurrent centroid sits slightly closer to culture than the primary centroid does (" + f"{tj['pc1_recurrent_mean']:.1f} against {tj['pc1_primary_mean']:.1f}, culture at {tj['pc1_culture']:.1f}" + "): do not say recurrence moves away from the dish."])

    # ---- 9 PCA
    s = content_slide(prs, "Primary and recurrent tumors separate, modestly")
    pvn = F["permanova"]
    fig_and_text(s, FIG / "chart_pca.png", [
        f"Each point is one tumor, placed by its 500 most variable genes. Every recurrent tumor lies to the right of every primary on the first axis, which carries {F['pc1']} % of the variance in those genes; the second carries {F['pc2']} %.",
        f"PERMANOVA on all genes puts {100 * pvn['r2']:.1f} % of the variance between the arms (F = {pvn['F']}), against 20 % expected for a random split of six tumors into three and three.",
        f"p = {pvn['p']:.2f} is the smallest value this test can return: with three per arm there are only ten ways to split six samples, and the observed split is the best-separated of them.",
        "A modest, directional separation that three per arm cannot push below p = 0.10; said out loud rather than left in the supplement."],
        img_w=5200000, size=14.5, gap=10, text_top=80000)
    refs(s, C.line("Principal components of the 500 most variable rlog genes; PERMANOVA (adonis2) on all genes. " + DESIGN))
    notes(s, ["1:48–2:03  Global structure. The arms are ordered on the first axis. Say the PERMANOVA p = 0.10 out loud. Frame it as: the arms are the best-separated split the design allows, and three per arm cannot return a smaller p. The expected R² for a random 3-vs-3 split is 1/5."])

    # ---- 10 only 35 genes move
    s = content_slide(prs, "Only 35 genes move, and the biggest movers are readthrough and non-coding loci")
    img = FIG / "chart_volcano.png"
    w, h = fit(img, 6000000, BOTTOM - TOP)
    picture(s, img, L, TOP, w, h)
    x2 = L + w + GAP
    tb = FIG / "table_fold.png"
    tw_, th_ = fit(tb, R - x2, 1650000)
    picture(s, tb, x2, TOP, tw_, th_)
    sw = FIG / "chart_threshold_sweep.png"
    sw_w, sw_h = fit(sw, R - x2, 2150000)
    picture(s, sw, x2, TOP + th_ + 100000, sw_w, sw_h)
    tot = [u + d for u, d in zip(thr["up"], thr["down"])]
    add_text(s, x2, TOP + th_ + 100000 + sw_h + 60000, R - x2, BOTTOM - (TOP + th_ + 100000 + sw_h + 60000), [
        f"{F['de_total']} genes pass FDR < 0.05 at two-fold ({F['de_up']} up, {F['de_down']} down); {tot[0]} at 1.5-fold, {tot[-1]} at four-fold. The six largest changes are three readthrough transcripts, a U6 snRNA and two unannotated loci, on 14 to 111 mean counts; five of the six are zero in every sample of one arm."],
        size=12, color=GREY, gap=0)
    refs(s, C.line(DESIGN + " Wald test with ashr shrinkage; FDR < 0.05 at every threshold; the sweep is Online Resource 1, S4."))
    notes(s, ["2:03–2:18  Differential expression. Only 35 genes. Walk the table once, do not read it: the six biggest fold changes are three readthrough transcripts, a U6 snRNA and two unannotated loci, on 14 to 111 mean counts against a median of 232 for the 35. Five of the six are all-or-none: every sample of one arm has no reads. TMEM189-UBE2V1 has 161 to 310 reads in each primary and none in any recurrence; RNU6-9 has none in any primary and 43 to 238 in each recurrence. They are large because one arm is empty. That is what justifies the pivot to gene sets, and it pre-empts the reviewer who asks why you did not chase the top hit.",
              f"Bottom right: the count moves with the fold-change threshold ({' / '.join(str(v) for v in tot)} genes at 1.5, 2, 2.83 and 4-fold) and roughly two in three are up at every threshold. The protein-coding changes cluster in matrix and calcium genes, taken up on the network slide."])

    # ---- 11 the translation machinery is turned down (running sums)
    s = content_slide(prs, "One dominant theme: translation is turned down")
    ti = rs[K["init"]]
    b10 = lambda k: f"{100 * rs[K[k]]['share_in_bottom_tenth']:.0f} %"  # noqa: E731
    s2m = cc["sets"][K["init"]]["manuscript_s2"]["mean_lfc"]
    fig_top(s, FIG / "chart_running_sum.png", [
        f"Of {F['gsea_n_sets']:,} gene sets, the six most depleted in recurrence are all translation or amino-acid stress. They are nominally significant (p ≤ {sets_p['p_max']:.3f}) and point the same way. Translation initiation leads (NES {m(top[0]['nes'])}, q = {top[0]['q']:.3f} in the published run), but its FDR q moves with the permutation seed: {min(q_full):.3f} to {max(q_full):.2f} over five seeds of the same analysis.",
        f"All {ti['n']} translation-initiation genes sit in the bottom {100 * ti['bottom_share']:.1f} % of the {ti['n_ranked']:,} ranked genes, each lower in recurrence (on average by about {100 * (1 - 2 ** s2m):.0f} %). In the bottom tenth sit {b10('elong')} of elongation and {b10('rib')} of ribosome members."],
        reserve=1150000, size=13, gap=5)
    refs(s, C.line(DESIGN + " GSEA running score (Diff_of_Classes ranking, gene-set permutation) from the pipeline's tables; fold changes, Online Resource 1, S2."))
    notes(s, ["2:18–2:43  The core result. Explain the plot in one breath: genes ranked from most up in recurrence to most down; walking down the list the score falls a little at every gene outside the set and jumps at every member, so a set whose members crowd the bottom drives the score to its floor just before the end. Then point at the ticks of translation initiation: all 80 at the far right.",
              f"Say it this way: all six are nominally significant and point the same way; which one clears FDR is not stable (translation initiation's q is below 0.05 at {n_full_ok} of 5 permutation seeds with all six tumors)." + " The other sets have a few members elsewhere: say 'crowd', not 'all'. None of these genes is significant alone; it is the coordinated shift that matters. If asked how robust it is: " + f"with any one tumor left out, translation initiation keeps nominal p < 0.001 and NES {m(max(ti_nes))} to {m(min(ti_nes))} in all 41 runs; without IL68B its q stays below 0.05 at every seed; without {', '.join(never)} it is above 0.05 at every seed (slide 15).",
              "If asked about the ranking metric: Diff_of_Classes is a linear difference of means, so abundant genes sit at the ends of the list; the manuscript's ashr-shrunk DESeq2 fold changes, on the log scale, agree in direction for every translation-initiation and ribosome gene."])

    # ---- 12 beyond translation
    s = content_slide(prs, "Beyond translation: growth and glycolysis fall, respiration does not agree")
    g = lambda k: themes[k]  # noqa: E731
    gly = [g(k) for k in ("HALLMARK_GLYCOLYSIS", "REACTOME_GLYCOLYSIS", "KEGG_GLYCOLYSIS_GLUCONEOGENESIS")]
    grow = [g(k) for k in ("GOBP_RIBOSOME_BIOGENESIS", "HALLMARK_MTORC1_SIGNALING", "HALLMARK_MYC_TARGETS_V1")]
    ox, ko, etc, tca, chol = (g("HALLMARK_OXIDATIVE_PHOSPHORYLATION"), g("KEGG_OXIDATIVE_PHOSPHORYLATION"), g("REACTOME_RESPIRATORY_ELECTRON_TRANSPORT"),
                              g("REACTOME_CITRIC_ACID_CYCLE_TCA_CYCLE"), g("REACTOME_CHOLESTEROL_BIOSYNTHESIS"))
    qmin = min(r["q"] for r in themes.values())
    fig_top(s, FIG / "chart_themes.png", [
        f"Growth signalling points the same way (ribosome biogenesis {m(grow[0]['nes'])}, mTORC1 {m(grow[1]['nes'])}, MYC {m(grow[2]['nes'])}; p {pv(min(r['p'] for r in grow))} to {pv(max(r['p'] for r in grow))}). Glycolysis falls in all three collections ({m(max(r['nes'] for r in gly))} to {m(min(r['nes'] for r in gly))}; p {pv(min(r['p'] for r in gly))} to {pv(max(r['p'] for r in gly))}).",
        f"Respiration does not agree across collections (Hallmark OXPHOS {m(ox['nes'])}, KEGG OXPHOS {m(ko['nes'], True)}, Reactome electron transport {m(etc['nes'], True)}), the TCA cycle does not move ({m(tca['nes'])}), and cholesterol synthesis rises ({m(chol['nes'], True)}, p = {pv(chol['p'])}). Nothing on this chart clears FDR (q ≥ {qmin:.2f})."],
        reserve=1000000, size=13, gap=5)
    refs(s, C.line(DESIGN + " Broad GSEA, gene-set permutation; Hallmark, KEGG, Reactome and GO collections as named; nominal p."))
    notes(s, ["2:43–2:55  About 12 s. The manuscript's line is 'the biosynthetic and respiratory machinery, not core carbon flux'. The respiration half depends on the collection: Hallmark OXPHOS is down, the KEGG, Reactome and GO respiration sets are up, and mitochondrial translation is −1.43 in GO but +1.54 in Reactome. So say 'glycolysis', not 'energy'. Cholesterol synthesis rising means not every anabolic program falls."])

    # ---- 13 one-sided at FDR
    s = content_slide(prs, "The effect is one-sided at q < 0.25")
    fig_and_text(s, FIG / "chart_gsea_landscape.png", [
        f"Every one of the {F['gsea_n_sets']:,} gene sets is a dot: its enrichment score against how far it clears FDR.",
        "In the published run six reach GSEA's exploratory q < 0.25, one of them q < 0.05; all six are down-regulated translation or stress sets. Those counts move with the permutation seed (slide 11).",
        f"Nothing on the up-regulated side survives correction (best q = {F['gsea_up_best_q']}); at nominal p the up side is led by proliferation (next slide).",
        "Past that threshold, the recurrent state is defined by what it turns down."],
        img_w=6500000, size=15, gap=10, text_top=80000)
    refs(s, C.line(DESIGN + " Broad GSEA, gene-set permutation, 1,000 permutations."))
    notes(s, ["2:55–3:08  One-sidedness. 8,869 sets tested, six clear q < 0.25, all on the same side and all the same theme. Nothing on the up side survives FDR; at nominal p the up side is not empty (406 sets up against 303 down at p < 0.05) and it is led by proliferation, which is the next slide."])

    # ---- 14 what else moves
    s = content_slide(prs, "What else moves: hypoxia and iron down, cell cycle up, all nominal")
    t_ = lambda k: TH[k]  # noqa: E731
    q2 = min(TH[k]["q"] for k in ("HIF1 targets", "hypoxia metagene", "hypoxia", "angiogenesis", "iron uptake and transport", "senescence", "EMT",
                                  "mitotic spindle", "G2M checkpoint", "E2F targets"))
    fig_and_text(s, FIG / "chart_themes2.png", [
        f"Hypoxic signalling falls, nominally: HIF1 targets and the Buffa hypoxia metagene both {m(t_('HIF1 targets')['nes'])} (p = {pv(t_('HIF1 targets')['p'])} and {pv(t_('hypoxia metagene')['p'])}), angiogenesis {m(t_('angiogenesis')['nes'])} (p = {pv(t_('angiogenesis')['p'])}). The simplest reading is that ablation removed the hypoxic core and what regrows is better perfused; perfusion was not measured.",
        f"Iron uptake and transport falls ({m(t_('iron uptake and transport')['nes'])}, p = {pv(t_('iron uptake and transport')['p'])}), as does senescence ({m(t_('senescence')['nes'])}, p = {pv(t_('senescence')['p'])}).",
        f"Cell-cycle sets rise: mitotic spindle {m(t_('mitotic spindle')['nes'], True)} (p = {pv(t_('mitotic spindle')['p'])}), G2M checkpoint {m(t_('G2M checkpoint')['nes'], True)}, E2F targets {m(t_('E2F targets')['nes'], True)}. A regrowing tumor dividing again, as a direction.",
        f"None of these clears FDR (q ≥ {q2:.2f}). The loss of hypoxic signalling, with the translational fall, is what a thermally stressed, reperfused margin would be expected to show: a check on the pipeline that also ranks the drugs."],
        img_w=6300000, size=13, gap=7, text_top=40000)
    refs(s, C.line(DESIGN + " Broad GSEA on the Semenza, Buffa, Hallmark, Reactome and Fridman sets named; nominal p, gene-set permutation."))
    notes(s, ["3:08–3:20  Cut if running long. Hypoxia is the cleanest story: HIF1 targets and the Buffa metagene both at −1.79 with nominal p near 0.01, angiogenesis down with them; the simplest reading is that the ablated core was the hypoxic core. Iron uptake down is worth one sentence because the top drug is an iron chelator, but do not build on it. The cell-cycle sets are up only nominally: say 'dividing again, as a direction'. The senescence and cell-cycle sets are not part of the positive-control argument."])

    # ---- 15 is the fall rat contamination, or one tumor?
    s = content_slide(prs, "The translational fall: rat contamination, or one tumor?")
    ms = mag["sets"]; lo = mag["leave_one_out"]["sets"]
    KS = ("initiation", "elongation", "ribosome")
    obs = [ms[k]["observed_mean_lfc"] for k in KS]
    mix = [ms[k]["mixture_bound_mean_lfc"] for k in KS]
    mat = [ms[k]["matched_background_mean_lfc"] for k in KS]
    k2 = [cc["sets"][K[k]]["k2"]["mean_lfc"] for k in ("init", "elong", "rib")]
    ratio_t = [ms[k]["median_ratio"] for k in KS]
    w68 = [lo[k]["without_IL68B"] for k in KS]
    m68 = [lo[k]["matched_without_IL68B"] for k in KS]
    others = [v for k in KS for kk, v in lo[k].items() if kk.startswith("without_") and kk != "without_IL68B"]
    assert lo["initiation"]["highest_tumour_counts"] == {"IL68B": lo["initiation"]["n"]} and lo["ribosome"]["highest_tumour_counts"] == {"IL68B": lo["ribosome"]["n"]}
    tw = cc["transcriptome"]; conc = {r["setting"]: r for r in cc["concordance"]}
    rng = lambda v: f"{m(max(v))} to {m(min(v))}" if f"{max(v):.2f}" != f"{min(v):.2f}" else m(v[0])  # noqa: E731
    fig_and_text(s, FIG / "chart_contamination_magnitude.png", [
        f"Recurrent tumors are {h_rec:.0f} % human against {h_pri:.0f} % for the primaries. Rat reads pull each gene toward its level in rat brain; across the transcriptome that pull is small (the {tw['n_detected']:,} genes detected in rat brain average {m(tw['mean_lfc_detected'], True)}, the rest {m(tw['mean_lfc_clean'], True)}).",
        f"Relative to their level in the tumors, translation genes are no more abundant in rat brain than a typical gene (control/tumor ratio {min(ratio_t):.2f} against {mag['median_ratio_tested']:.2f}). Counting every shared read as rat, contamination could move them by {m(min(mix))} at most; they fall {rng(obs)}, against {rng(mat)} for matched genes.",
        f"Most of the fall comes from one primary tumor: IL68B is the highest of the six on every translation-initiation and ribosome gene. Without it the sets fall {rng(w68)} (matched genes {rng(m68)}); dropping any other tumor leaves {m(max(others))} to {m(min(others))}. The enrichment does not depend on it: without IL68B, translation initiation stays at q < 0.05 at all five permutation seeds.",
        [(f"Adjusting for factors estimated from the controls{C('risso')} leaves {rng(k2)} (k = 2), but those factors also separate the arms. Purity, one tumor and arm cannot be pulled apart in six animals.", BLUE)]],
        img_w=5900000, size=12, gap=6, text_top=20000)
    refs(s, C.line("Repository analysis, not in the manuscript. Bound: every shared read treated as rat, with the controls' profile. Matched: 25 nearest non-translation genes in abundance and control/tumor ratio. Without IL68B: median-of-ratios fold changes of the other five."))
    notes(s, ["3:20–3:42  About 22 s: the most important caveat of the talk, and the slide most likely to draw a question. Two questions, two answers.",
              "Contamination: every translation gene is detected in the rat-brain controls, as most abundant conserved genes are, but relative to their tumor level they are no more abundant in rat brain than a typical gene, so contamination predicts a shift of about zero for them. They fall 0.40 on the log2 scale; genes matched on abundance and that ratio fall about a third as much.",
              f"One tumor: IL68B, a primary, is the highest of the six on every translation-initiation and ribosome gene while genome-wide it is average. Without it the sets fall {rng(w68)}, a quarter of the size, still more than matched genes ({rng(m68)}); dropping any other tumor makes the fall larger. The gene-set test was re-run with each tumor left out (41 runs of the published command, five seeds): " + f"without IL68B translation initiation stays at q {min(q_68):.3f} to {max(q_68):.3f}; it is leaving out {', '.join(never)} that pushes q above 0.05. Nominal p stays below 0.001 in every run. Say plainly: IL68B carries the size of the fall, not its rank.",
              f"Adjustment: the RUVSeq factors estimated from the controls are higher in every primary than in every recurrent tumor, so adjusting removes part of the arm difference itself; at k = 1 the separation between arms collapses (silhouette {m(conc['k1']['pri_rec_silhouette'])}). That is collinearity, not a refutation. This analysis is not in the manuscript.",
              "The unadjusted run here is the RUVSeq script's own DESeq2 without shrinkage, so its fold changes are larger than the manuscript's ashr-shrunk ones; the adjusted runs test 19,385 genes and the means are taken over the unadjusted run's 14,849."])

    # ---- 16 subtypes
    s = content_slide(prs, "Scored against the published classifiers: the astrocyte-like score falls")
    gb = F["gsva_best"]; ac, mt = sub["Neftel_AC"], sub["Garofano_MTC"]
    fig_and_text(s, FIG / "chart_subtypes.png", [
        f"Each tumor is scored by GSVA for how much of each published signature it expresses: ten Neftel{C('neftel')} and Garofano{C('garofano')} signatures, 39 to 50 genes each as published.",
        f"One transition reaches nominal significance: the astrocyte-like score falls, {m(gb['primary'], True)} → {m(gb['recurrent'], True)} (p = {ac['p']:.3f}; q = {ac['q_bh']:.3f} across the ten).",
        f"The mitochondrial signature falls too ({m(mt['change'])}, p = {mt['p']:.2f}, not significant).",
        "Mesenchymal and neural-progenitor scores drift but do not reach significance."],
        img_w=6100000, size=14, gap=9, text_top=60000)
    refs(s, C.line(DESIGN + " GSVA per tumor; Benjamini–Hochberg across the ten signatures."))
    notes(s, ["3:42–3:57  Subtypes. Ten published signatures, scored per sample by GSVA over the published gene lists; using the published definitions rather than a hand-picked panel is the defence. One transition is nominally significant: a fall in the astrocyte-like score, p = 0.007, q = 0.075 after correction across the ten, so none clears FDR.",
              f"Mitochondrial signature: −0.58, p = {mt['p']:.2f}, q = {mt['q_bh']:.2f}. It shares genes and tumors with the Hallmark OXPHOS set, so the two agree without confirming each other. Note that the supplement's adjusted-p column is mis-ordered (it prints 0.075 for the mitochondrial signature); the q values on this chart are recomputed."])

    # ---- 17 hubs
    s = content_slide(prs, "The hubs: matrix mostly down, calcium up")
    img = FIG / "chart_ppi.png"
    tb = FIG / "table_hubs.png"
    tw_, th_ = fit(tb, 4600000, BOTTOM - TOP)
    picture(s, tb, R - tw_, TOP, tw_, th_)
    cw = R - tw_ - GAP - L
    w, h = fit(img, cw, BOTTOM - TOP - 1850000)
    picture(s, img, L + (cw - w) // 2, TOP, w, h)
    add_text(s, L, TOP + h + 60000, cw, 420000,
             [f"{len(ppi['degree'])} of the 35 genes have a STRING interaction. Navy down, red up; filled, degree ≥ 3; line weight, STRING score."],
             size=CAP, color=GREY, gap=0)
    add_text(s, L, TOP + h + 520000, cw, BOTTOM - (TOP + h + 520000), [
        "Down: most of the structural matrix, type I collagen, biglycan, prolargin and hemicentin; collagen XVII and glypican 3 rise. Up: neuronal calcium handling and a presynaptic scaffold, calbindin, ryanodine receptor 2 and piccolo.",
        f"Whether that echoes the neuron–glioma synapse described in high-grade glioma{C('venkatesh', 'venkataramani')}, or only reflects which cells survived, six animals cannot tell us."],
        size=12.5, gap=4)
    refs(s, C.line("STRING interactions among the 35 differentially expressed genes. " + DESIGN))
    notes(s, ["3:57–4:08  Cut if running long. Do not read the table; point at the sign column. The fibrillar matrix genes (COL1A1, BGN, PRELP, HMCN1) go down with IGFBP3 and LRP1; COL17A1 and GPC3 go up. The calcium and presynaptic genes (CALB1, RYR2, PCLO) go up. Offer the neuron–glioma synapse as a question, not a claim, and say plainly that six animals cannot separate it from a change in which cells survived. Know the two Nature 2019 papers."])

    # ---- 18 the drug filter
    s = content_slide(prs, f"From {fun['n_compounds']} opposing compounds to {fun['n_both_bbb']} that clear both barrier models")
    fig_and_text(s, FIG / "chart_funnel.png", [
        f"DSigDB holds a gene set for each drug.{C('yoo')} The pipeline asks which drug's set sits among the genes that fall in recurrence, then weights by predicted brain penetration (ADMET-AI{C('swanson')}), with a second barrier model as a check (BOILED-Egg{C('daina')}).",
        f"The {fun['n_signatures']} most strongly opposing drug gene sets name {fun['n_compounds']} compounds. {fun['n_clinical']} have a clinical phase and a structure the barrier model can score; the rest include phosphine, tributyltin and a flame retardant.",
        f"{fun['n_both_bbb']} clear both barrier models, and ciclopirox ranks first among them and overall.",
        "The filter's limits, said up front: it keeps things that are not systemic therapies (oxygen, ozone, magnesium) and drops approved drugs with no recorded phase, deferoxamine among them."],
        img_w=5200000, size=13.5, gap=8, text_top=40000)
    refs(s, C.line("Score: |NES| to the power 1.5, weighted by predicted blood–brain barrier permeability; Online Resource 1, S12."))
    notes(s, ["4:08–4:23  Therapeutics. Explain the score in one sentence: how strongly the drug's gene set sits among the genes that fall in recurrence, weighted by predicted barrier penetration. Then the funnel: " + f"{fun['n_compounds']} compounds, {fun['n_clinical']} with a phase, {fun['n_both_bbb']} that clear both barrier models" + ". Cobalt appears twice in the supplement, under two capitalisations, and is counted once. Own the filter's limits before anyone asks. The 100 is a cap on the most opposing signatures, not the total that opposed."])

    # ---- 19 where every clinically available candidate sits
    s = content_slide(prs, "Where every clinically available candidate sits")
    strong = [x for x in ds["strongest_overall"]][:2]
    assert {x["drug"].lower() for x in strong} == {"dimethyloxalylglycine", "deferoxamine"} and not any(x["clinic"] for x in strong), strong
    fig_and_text(s, FIG / "chart_drug_scatter.png", [
        f"Each point is one of the {ds['n_with_phase']} compounds with a clinical phase: how strongly its gene set opposes the recurrence signature, against its predicted barrier permeability; size is the combined score.",
        "Ciclopirox is alone at the top right. Filled points are those both barrier models predict will cross.",
        f"The two strongest reversals of all, {strong[0]['drug'].lower()} and {strong[1]['drug'].lower()} (|NES| {strong[0]['absnes']:.2f} each), have no recorded clinical phase and are not on this chart; deferoxamine is an approved iron chelator that the phase filter drops.",
        "Just left of ciclopirox at the top sit pentetrazol, primidone and nilutamide; the prior-art audit sets them aside."],
        img_w=6500000, size=13.5, gap=8, text_top=60000)
    refs(s, C.line("Online Resource 1, S12: ADMET-AI barrier probability, BOILED-Egg agreement. " + DESIGN))
    notes(s, ["4:23–4:31  Cut if running long. A redrawn version of Figure 2A: only the 54 compounds with a phase, on the ADMET-AI probability, with the two barrier models marked. The two axes are the whole of the score. The manuscript text names LY-294002 among the strongest reversals; by |NES| it is sixth, below ciclopirox and metformin."])

    # ---- 20 the candidate
    s = content_slide(prs, "Ciclopirox ranks first with or without the barrier weight")
    img = FIG / "chart_drugs.png"
    tb = FIG / "table_drugs.png"
    tw_, th_ = fit(tb, 4900000, 2500000)
    picture(s, tb, R - tw_, TOP, tw_, th_)
    w, h = fit(img, R - tw_ - GAP - L, BOTTOM - TOP)
    picture(s, img, L, TOP, w, h)
    cpx = s13[0]; nxt = [x for x in s13[1:] if x["both_agree"]][0]; run = s13[1]
    assert cpx["drug"] == "ciclopirox" and cpx["rank"] == 1 and cpx["rank_unweighted"] == 1
    assert run["rank"] == 2 and not run["both_agree"] and run["egg"] == "out", run
    add_text(s, R - tw_, TOP + th_ + 120000, tw_, BOTTOM - (TOP + th_ + 120000), [
        f"Ciclopirox scores {cpx['score']:.2f} against {nxt['score']:.2f} for {nxt['drug']}, the next compound on which both barrier models agree; the overall runner-up, {run['drug']} ({run['score']:.2f}), fails the second model. It ranks first with or without the permeability weight.",
        f"It is approved (a topical antifungal), and an oral form has completed phase 1 in patients with hematologic cancers.{C('minden')}",
        [("Computational predictions: no agent named here has been tested in this model or is approved for this use.", BLUE)]],
        size=12.5, gap=6)
    refs(s, C.line("Rank: position among the 54 clinically available compounds by the BBB-weighted score (Online Resource 1, S12–S13)."))
    notes(s, ["4:31–4:46  The candidate. Rank column: overall rank among the 54. Ciclopirox is approved, ranked first with and without the permeability weight, and both barrier models agree. Deliver the caveat sentence exactly as written; this is where a reviewer will push."])

    # ---- 21 the prior-art audit
    s = content_slide(prs, "What is already known about the top twenty")
    tA = [n for _, n in sorted(tiers["A"])]
    fig_and_text(s, FIG / "table_prior_art.png", [
        f"The twenty highest-ranked of the {ds['n_with_phase']} compounds were read against the glioma literature. Tier A: in vivo or clinical glioma evidence. B: in vitro, or in vivo but contested. C: none, failed or refuted, or not a therapy at all.",
        f"Ciclopirox is tier A: it inhibits U251 and three other lines in vitro and a subcutaneous xenograft{C('su')}, and an oral form completed phase 1.{C('minden')} The other As are ifosfamide, an alkylator the screen was bound to recover, thioridazine and ganciclovir.",
        f"Seven have no glioma literature; primidone was tested and failed; pentetrazol is a convulsant; amiodarone use was linked to worse glioblastoma survival{C('boursi')} and diazepam antagonises temozolomide in vitro.{C('drljaca')}",
        [("An audit of what has been published, not a test of anything.", BLUE)]],
        img_w=5800000, size=12, gap=6, text_top=20000)
    refs(s, C.line("Online Resource 1, S13 and S15; PubMed searched per compound."))
    tier_rank = {n: r for t_k in tiers for r, n in tiers[t_k]}
    notes(s, [f"4:46–4:54  Cut if running long. Point at the tier column: ciclopirox is the only A in the top five, and ranks 2 to 5 are all C. The other As are ifosfamide ({tier_rank['ifosfamide']}), thioridazine ({tier_rank['thioridazine']}) and ganciclovir ({tier_rank['ganciclovir']}, tested only as an HSV-tk prodrug). {len(tiers.get('B', []))} are B and {len(tiers.get('C', []))} are C. The audit is what separates a ranked list from a shortlist.",
              "Amiodarone: in 1,076 glioblastoma patients, active users had worse overall survival (HR 4.41). Diazepam: reduces temozolomide-induced apoptosis in U87 cells. The supplement's prior-art sheet marks thioridazine as outside the ranked list; it is rank 12 and is shown here."])
    assert tA[0] == "ciclopirox"

    # ---- 22 why ciclopirox could work
    s = content_slide(prs, "Why ciclopirox: the tumor depends on the axis it hits")
    e = lambda gname: float(dm[gname]["u251_gene_effect"])  # noqa: E731
    md = lambda gname: float(dm[gname]["median_across_lines"])  # noqa: E731
    assert all(e(x) < -0.5 for x in ("DOHH", "DHPS", "EIF5A")) and md("DOHH") > -0.5 and md("DHPS") < -0.5 and md("EIF5A") < -0.5
    assert e("HIF1A") > 0 and -0.5 < e("TFRC") and md("TFRC") < -0.5 and dm["RRM1"]["common_essential"] == "True" and dm["RRM2"]["common_essential"] == "True"
    fig_and_text(s, FIG / "chart_depmap.png", [
        "Ciclopirox chelates iron and inhibits deoxyhypusine hydroxylase (DOHH), the enzyme that activates eIF5A, a factor the ribosome needs to keep elongating.",
        f"In DepMap's genome-wide CRISPR screen of U-251 MG in culture{C('depmap')}, knocking out DOHH ({m(e('DOHH'))}), DHPS ({m(e('DHPS'))}) or EIF5A ({m(e('EIF5A'))}) costs fitness past the −0.5 dependency threshold. Most lines need DHPS and EIF5A too (median {m(md('DHPS'))} and {m(md('EIF5A'))}); DOHH is the step the typical line does not need (median {m(md('DOHH'))}).",
        f"Ribonucleotide reductase, the other target, is needed by nearly every line (U-251 RRM1 {m(e('RRM1'))}, RRM2 {m(e('RRM2'))}). HIF1A is not a dependency ({m(e('HIF1A'), True)}), nor, unusually, is the transferrin receptor ({m(e('TFRC'))} against a median of {m(md('TFRC'))}).",
        f"Delivery: LITT opens the blood–brain barrier at the margin for weeks after ablation{C('cleary')}, which is the window.",
        [("Public screen data and a computational prediction; nothing has been dosed in this model.", BLUE)]],
        img_w=6100000, size=12.5, gap=6, text_top=20000)
    refs(s, C.line())
    notes(s, ["4:54–5:14  About 20 s. The answer to 'why should an iron chelator work on a cell that is building less': ciclopirox also inhibits DOHH, and U-251 depends on the whole hypusination axis. DHPS and EIF5A are needed by most lines, so they are not a selective weakness; DOHH is the step most lines can do without, and U-251 needs it more than the typical line (" + f"{m(e('DOHH'))} against {m(md('DOHH'))}" + "), a modest margin. Say plainly that this is a public screen of the line in culture, not our data, and that nothing has been dosed. The barrier window is the delivery argument."])

    # ---- 23 conclusions
    s = content_slide(prs, "Conclusions")
    img = FIG / "chart_arms.png"
    w, h = fit(img, 5500000, BOTTOM - TOP - 1000000)
    picture(s, img, L, TOP, w, h)
    x2 = L + w + GAP
    hy, mg = TH["HIF1 targets"], TH["hypoxia metagene"]
    add_text(s, x2, TOP, R - x2, h, [
        f"Post-LITT recurrence turns its translation machinery down: initiation (nominal p < 0.001 whichever tumor is left out; FDR q {min(q_full):.2f} to {max(q_full):.2f} across permutation seeds), elongation and the ribosome, with the GCN2 amino-acid set (mostly ribosomal proteins) and growth signalling lower alongside, nominally. One primary tumor carries most of the size of the difference, not its rank.",
        "Glycolysis falls in every collection; respiration does not agree across collections, and the TCA cycle is untouched.",
        f"Hypoxic signalling falls, nominally (HIF1 targets and hypoxia metagene both {m(hy['nes'])}; q = {hy['q']:.1f}): the simplest reading is that ablation removed the hypoxic core.",
        f"Of ten published subtype signatures, only the astrocyte-like score falls at nominal significance (p = {ac['p']:.3f}; q = {ac['q_bh']:.3f}). Ciclopirox ranks first with and without the permeability weight, is approved as a topical antifungal, completed phase 1 in cancer{C('minden')}, and was nominated independently by another group's signature-reversal screen.{C('sun')}"],
        size=12.5, gap=5)
    navy_bar(s, L, TOP + h + 120000, R - L, BOTTOM - (TOP + h + 120000),
             "These are the changes expected of a thermally stressed, reperfused margin; recovering them from an unbiased transcriptome is a check on the pipeline that also ranks the drugs.", size=15)
    refs(s, C.line("Evidence status: investigational; one cell line (U251N) in one rat xenograft model with MRI-guided LITT, n = 3 per arm; drug candidates are computational predictions; GEO GSE338105."))
    notes(s, ["5:14–5:36  Conclusions. Four beats, then the closing line. The closing argument is that the translational, growth, glycolytic and hypoxic arms move together, not any single q value: only translation initiation clears FDR. The other group's screen is a signature-reversal analysis of glioblastoma that also nominated ciclopirox; say 'nominated', not 'validated'. Say the closing sentence and stop."])

    # ---- 24 limitations and next steps
    s = content_slide(prs, "Limitations and next steps")
    half = (R - L - GAP) // 2
    left = ["Limitations",
            "n = 3 per arm. Gene-set results are best read as directional, which is why nominal p is reported alongside FDR q.",
            "A single cell line (U251N) in a single xenograft model: no patient tissue, and no T-cell immunity in the athymic rat.",
            f"Bulk RNA-seq averages the margin with whatever else was harvested, and purity differs by arm ({h_pri:.0f} % against {h_rec:.0f} % human); the persister state is inferred, not isolated.",
            f"The size of the translational fall leans on one primary tumor (without it, a quarter as large); its FDR q moves with the permutation seed ({min(q_full):.2f} to {max(q_full):.2f}) and with which tumor is left out.",
            "A methylation array on the same six tumors (866,238 probes) found no probe past FDR (smallest q = 0.09); with purity and arm confounded it cannot separate the two.",
            "Every drug candidate is a computational prediction; none has been tested in this model."]
    right = ["Next steps",
             "Spatial or single-cell profiling of the sublethal margin, to test whether the persister state sits where we think it does.",
             "Puromycin incorporation and polysome profiling, to confirm that translation is suppressed and not only transcriptionally lower.",
             "Ciclopirox dosed into the post-ablation barrier window, with survival as the endpoint.",
             "Validation in an immunocompetent model, and in patient tissue from the peri-ablation zone."]
    for x, block in ((L, left), (L + half + GAP, right)):
        head_box(s, x, TOP, half, 460000, block[0], [], size=22)
        add_text(s, x, TOP + 620000, half, BOTTOM - TOP - 620000, block[1:], size=13.5, gap=7)
    refs(s, C.line(DESIGN + " Methylation: Illumina EPIC on the same tumors, primary vs recurrent, BH over 866,238 probes; purity from the matched RNA-seq graft fraction."))
    notes(s, ["5:36–5:48  Keep if at all possible: n, purity and the computational caveat are on the results slides, but the methylation result and the next steps are said only here. Anticipated questions: why not more animals; why gene-set permutation; has any drug been tested (no); is the fall contamination (slide 15).",
              "The methylation line is from REVIEW/purity_confound.py run on beta_values.rda: 0 probes at q < 0.05, 39 at q < 0.25, smallest q 0.092; the r = −0.92 between effect size and purity across the 9,079 nominally changed probes is NOT said: any covariate correlated with arm as strongly returns about the same value, so it says nothing about purity specifically. Not in either manuscript."])

    # ---- 25 acknowledgments and data availability
    s = content_slide(prs, "Acknowledgments and data availability")
    add_text(s, L, TOP, 6300000, BOTTOM - TOP, [
        "Tavarekere N. Nagaraja, PhD, Indrani Datta, DHI, and Ian Y. Lee, MD, Hermelin Brain Tumor Center, Henry Ford Health, who built and characterized the model this analysis rests on.",
        "Supported by a Henry Ford Health Physician Scientist Award A20050 (I.Y.L.).",
        "Animal procedures were approved by the Henry Ford Health IACUC (protocol #1509) and conducted in accordance with the ARRIVE guidelines.",
        "Sequencing data: GEO GSE338105. Differential-expression, enrichment and drug-ranking tables are released as supplementary data."], size=15, gap=10)
    navy_bar(s, R - 3600000, TOP + 300000, 3600000, 1300000, ["Questions", "go2432@wayne.edu"], size=20)
    add_text(s, R - 3600000, TOP + 1800000, 3600000, 1200000,
             ["Sequencing data: GEO GSE338105", "No commercial support."],
             size=13, color=GREY, gap=6, align=PP_ALIGN.CENTER)
    notes(s, ["5:48–6:00  Thanks. Name the Henry Ford group. Leave the GEO accession on screen during questions."])
    assert C.line() == ""

    for sl in prs.slides:                                   # a number never parts from its '%'
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if " %" in r.text:
                            r.text = r.text.replace(" %", "\u00a0%")
    prs.save(str(out))
    print("wrote", out, "slides:", len(prs.slides))
    print("references by first appearance:")
    for k, n in sorted(C.num.items(), key=lambda t: t[1]):
        print(f"  {n:2d}. {REFS[k]}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE / "CNS2026_Schwing_Abstract418_CNStemplate.pptx"))
    a = ap.parse_args()
    if not TEMPLATE.exists():
        raise SystemExit(f"CNS speaker template missing: {TEMPLATE} (the CNS's file, gitignored; see SLIDES/templates/README.md)")
    build(Path(a.out))
