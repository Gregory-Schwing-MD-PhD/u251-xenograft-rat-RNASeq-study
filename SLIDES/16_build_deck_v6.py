# -*- coding: utf-8 -*-
"""The CNS 2026 Abstract 418 talk, version 6 (2026-09-27): one panel per slide, the panel large, the text short.

What v6 changes, all from Greg's review of v5 on 2026-09-27:
  1. The contralateral / operating-room slide stays IN the talk.
  2. More of the backup material is promoted into the talk, in a logical order.
  3. The talk is six minutes: the deck marks a 6-minute core and carries the rest as the expanded talk, with a cut
     order, rather than pretending every slide fits.
  4. Nothing says the instrument reading is what was measured. Every instrument sentence is conditional or
     interrogative; the two readings are given as alternatives this design cannot separate.
  5. No slide carries a multi-panel figure any more. Every results slide shows ONE panel, laid out by v6_layout so
     the panel takes the largest area the content box allows, with at most three short lines.

Build:
    python SLIDES/13_v5_figures.py        # slide-size panels (v5 set, still used for the single-axes figures)
    python SLIDES/15_slide_panels.py      # single-panel crops of every multi-panel figure -> figures_cns/v6/
    python SLIDES/09_build_deck_v4.py     # only if the v4 deck is missing
    python SLIDES/16_build_deck_v6.py

The builder types no result number: every value comes from figures_cns/v5_data (byte copies of the grid outputs),
and the load-bearing ones are asserted before a slide is made. It prints, for each slide, the share of the content
area its panel covers, and refuses to finish if a results slide falls below FLOOR.
"""
from __future__ import annotations

import argparse
import csv
import importlib.util
import json
import sys
from pathlib import Path

from pptx import Presentation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

_spec = importlib.util.spec_from_file_location("deck3_v6", HERE / "04_build_cns_template_deck.py")
d3 = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(d3)

from cns_template import BLUE, notes  # noqa: E402
from v6_layout import auto_slide, panel_fit, panel_slide, side_panel_slide, statement_slide  # noqa: E402

V4 = HERE / "CNS2026_Schwing_Abstract418_CNStemplate_v4.pptx"
FIG5 = HERE / "figures_cns" / "v5"
FIG6 = HERE / "figures_cns" / "v6"
DAT = HERE / "figures_cns" / "v5_data"
DESIGN = d3.DESIGN
# A results panel must either cover FLOOR of the content area or be limited by the slide itself (FILL of the
# dimension that binds). A portrait panel at full slide height covers only about 48 % of a 2.1:1 box and is still
# as large as it can be; a small panel floating in the middle is the defect v6 exists to fix.
FLOOR, FILL = 0.45, 0.97


def J(name):
    with open(DAT / name, encoding="utf-8") as fh:
        return json.load(fh)


def T(name):
    with open(DAT / name, encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def pct(x, nd=0):
    return f"{100 * float(x):.{nd}f}"


def sup(n):
    return str(n).translate(str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹"))


def fig(name):
    """A panel, preferring the single-panel crop in v6 over the whole figure in v5."""
    for d in (FIG6, FIG5):
        p = d / name
        if p.exists():
            return p
    raise SystemExit(f"missing panel {name}: run SLIDES/15_slide_panels.py and SLIDES/13_v5_figures.py first")


AREAS = []


def slide(kind, prs, title, image=None, lines=(), ref=None, note=None, core=False, seconds=None, **kw):
    """One slide, recorded for the build's self-check and for the timing table it prints."""
    tag = ("CORE " if core else "     ") + (f"{seconds:>3d}s " if seconds else "     ")
    if kind == "statement" and image is not None:
        raise SystemExit(f"{title!r}: a statement slide takes lines=..., not a positional list (it landed in `image`)")
    if kind != "statement" and not lines:
        raise SystemExit(f"{title!r}: a results slide with no text is almost always the same positional-argument slip")
    if kind == "statement":
        s = statement_slide(prs, title, list(lines), ref=ref, note=note, **kw)
        AREAS.append((tag, title, None, None, None))
        return s
    p = fig(image)
    share, fill = panel_fit(p, n_lines=list(lines))
    s = auto_slide(prs, title, p, lines=list(lines), ref=ref, note=note, **kw)
    AREAS.append((tag, title, image, share, fill))
    return s


def build(out: Path, order_module=None):
    if not V4.exists():
        raise SystemExit(f"missing {V4}: run python SLIDES/09_build_deck_v4.py first")
    prs = Presentation(str(V4))
    if len(prs.slides) != 21:
        raise SystemExit(f"v4 deck has {len(prs.slides)} slides, expected 21")
    if order_module is None:
        raise SystemExit("16_build_deck_v6.py: the running order is not wired in yet (see v6_order.py)")
    order_module.build_slides(prs, slide)
    order_module.arrange(prs)

    for sl in prs.slides:                       # a number never parts from its '%'
        for sh in sl.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if " %" in r.text:
                            r.text = r.text.replace(" %", " %")
    prs.save(str(out))

    print(f"wrote {out}  slides: {len(prs.slides)}")
    bad = []
    for tag, title, image, share, fill in AREAS:
        if share is None:
            print(f"  {tag} {title[:58]:58s} (no figure)")
            continue
        ok = share >= FLOOR or fill >= FILL
        print(f"  {tag} {title[:58]:58s} {image[:34]:34s} area {share:.0%} fill {fill:.0%}"
              f"{'' if ok else '   <-- TOO SMALL'}")
        if not ok:
            bad.append((title, image, share, fill))
    if bad:
        raise SystemExit(f"{len(bad)} slide(s) below the panel floor (area {FLOOR:.0%} or fill {FILL:.0%}): " +
                         "; ".join(f"{t} ({i}, area {s:.0%}, fill {f:.0%})" for t, i, s, f in bad))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE / "CNS2026_Schwing_Abstract418_CNStemplate_v6.pptx"))
    a = ap.parse_args()
    try:
        import v6_order
    except ModuleNotFoundError:
        v6_order = None
    build(Path(a.out), v6_order)
