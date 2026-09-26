"""Derive every number that appears on the CNS slides from the analysis outputs.

Nothing on a slide is hand-typed: 03_build_deck.py reads figures/facts.json.
Run order: 01_extract_facts -> 02_make_figures -> 03_build_deck
"""
import csv, json, os, re

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
FIG  = os.path.join(HERE, "figures")
GSEA = os.path.join(ROOT, "ANALYSIS/results_therapy_v3/report_gsea")
DE   = os.path.join(ROOT, "ANALYSIS/results_therapy_v3/tables/differential",
                    "therapy_impact.deseq2.results_filtered.tsv")
DRUG = os.path.join(ROOT, "manuscript/therapy_impact_Drug_Profiles_Comprehensive.csv")
LLM  = os.path.join(ROOT, "manuscript/therapy_impact_LLM_Analysis_Report.txt")

f = {}

# --- Broad GSEA, gene-set permutation (the test reported in the manuscript) ----
rows = []
for fn in sorted(os.listdir(GSEA)):
    if fn.endswith(".tsv"):
        rows += list(csv.DictReader(open(os.path.join(GSEA, fn), encoding="utf-8"),
                                    delimiter="\t"))
rows = [r for r in rows if r.get("NES") not in (None, "", "NA")]
for r in rows:
    r["NES"], r["p"], r["q"] = float(r["NES"]), float(r["NOM p-val"]), float(r["FDR q-val"])
rows.sort(key=lambda r: r["NES"])
f["gsea_n_sets"]  = len(rows)
f["gsea_down_top"] = [{"name": r["NAME"], "nes": round(r["NES"], 2),
                       "p": r["p"], "q": round(r["q"], 3), "size": int(r["SIZE"])}
                      for r in rows[:10]]
up = [r for r in rows if r["NES"] > 0]
f["gsea_up_best_q"]   = round(min(r["q"] for r in up), 2)
f["gsea_up_best_nes"] = round(max(r["NES"] for r in up), 2)
gsea_rows = rows          # `rows` is reused further down for the Excel sheets

# --- DESeq2 ------------------------------------------------------------------
de = list(csv.DictReader(open(DE, encoding="utf-8"), delimiter="\t"))
sig = [r for r in de if r["padj"] not in ("", "NA") and float(r["padj"]) < 0.05
       and abs(float(r["log2FoldChange"])) > 1]
f["de_up"]    = sum(1 for r in sig if float(r["log2FoldChange"]) > 0)
f["de_down"]  = sum(1 for r in sig if float(r["log2FoldChange"]) < 0)
f["de_total"] = len(sig)

# --- DSigDB drug ranking ------------------------------------------------------
dr = list(csv.DictReader(open(DRUG, encoding="utf-8")))
f["drug_top"] = [{"rank": int(r["Rank"]), "drug": r["Drug"], "nes": float(r["NES"]),
                  "score": float(r["IntegratedScore"]), "bbb": float(r["BBB_Score"])}
                 for r in dr[:10]]
f["drug_by_name"] = {}
for r in dr:
    f["drug_by_name"].setdefault(r["Drug"], {"rank": int(r["Rank"]),
                                             "score": float(r["IntegratedScore"]),
                                             "nes": float(r["NES"]),
                                             "bbb": float(r["BBB_Score"])})

# --- PCA / PERMANOVA / subtypes / PPI hubs (pipeline analysis report) ---------
txt = open(LLM, encoding="utf-8").read()
m = re.search(r"PCA variance explained: PC1=([\d.]+)%, PC2=([\d.]+)%", txt)
f["pc1"], f["pc2"] = float(m.group(1)), float(m.group(2))
m = re.search(r"PERMANOVA.*?R2=([\d.]+), F=([\d.]+), p=([\d.]+)", txt)
f["permanova"] = {"r2": float(m.group(1)), "F": float(m.group(2)), "p": float(m.group(3))}
f["subtypes"] = [{"name": n, "primary": float(a), "recurrent": float(b), "sig": bool(s)}
                 for n, a, b, s in re.findall(
                     r"^  (\S+)\s+Primary_U2=([-+][\d.]+) -> Recurrent_U2=([-+][\d.]+)"
                     r"(\s+\[sig[^\]]*\])?$", txt, re.M)]
f["ppi_hubs"] = re.search(r"PANEL E - PPI HUB GENES \(top by degree\): (.+)",
                          txt).group(1).split(", ")

# --- the STRING network, and which way each of its genes moved ----------------
import openpyxl
wb = openpyxl.load_workbook(os.path.join(ROOT, "manuscript",
                                         "Supplementary_Data.xlsx"), read_only=True)
lfc = {}
rows = wb["S2_DE_significant"].iter_rows(values_only=True)
head = next(rows)
si, li = head.index("symbol"), head.index("log2FoldChange")
for r in rows:
    if r[si]:
        lfc[str(r[si])] = float(r[li])
edges = list(wb["S6_PPI_edges"].iter_rows(values_only=True))[1:]
nodes = sorted({str(a) for a, b, _ in edges} | {str(b) for a, b, _ in edges})
f["ppi_nodes"] = {n: round(lfc[n], 2) for n in nodes if n in lfc}
f["ppi_nodes_down"] = sorted([n for n, v in f["ppi_nodes"].items() if v < 0],
                             key=lambda n: f["ppi_nodes"][n])
f["ppi_nodes_up"] = sorted([n for n, v in f["ppi_nodes"].items() if v > 0],
                           key=lambda n: -f["ppi_nodes"][n])
f["de_significant"] = sorted(
    ({"symbol": s, "lfc": round(v, 2)} for s, v in lfc.items()),
    key=lambda d: d["lfc"])

# --- 2026 reanalysis -------------------------------------------------------
# The subtype scores and the drug ranking were both rebuilt (see subtypes/).
# Subtypes: real GSVA over the published Neftel and Garofano signatures, in
# place of mean-z over hand-written 5-6 gene panels. Drugs: ADMET-AI BBB in
# place of a hand-rolled rubric, restricted to compounds that reached a clinic.
SUB = os.path.join(ROOT, "subtypes")

sub = list(csv.DictReader(open(os.path.join(SUB, "subtype_rerun_results.csv"),
                               encoding="utf-8")))
f["gsva"] = [{"name": r["Signature"],
              "primary": round(float(r["Mean primary"]), 3),
              "recurrent": round(float(r["Mean recurrent"]), 3),
              "change": round(float(r["Change"]), 3),
              "p": float(r["p"]), "q": float(r["q"])} for r in sub]
f["gsva_by_name"] = {r["name"]: r for r in f["gsva"]}
f["gsva_best"] = min(f["gsva"], key=lambda r: r["p"])

drg = list(csv.DictReader(open(os.path.join(SUB, "drug_ranking_final.csv"),
                               encoding="utf-8")))
f["drug_clinical_n"] = len(drg)
f["drug_tierA"] = [{"drug": r["Drug"], "phase": int(float(r["phase"])),
                    "nes": float(r["NES"]), "bbb": float(r["BBB_Martins"]),
                    "egg": r["BOILED_Egg"],
                    "score": round(float(r["score_bbb"]), 2)}
                   for r in drg if r["both_agree"].lower() == "true"]
f["drug_top"] = f["drug_tierA"][0] if f["drug_tierA"] else None

# thematic enrichment: pull the named sets the narrative rests on
_want = {
    "HALLMARK_MTORC1_SIGNALING": "mTORC1 signalling",
    "HALLMARK_MYC_TARGETS_V1": "MYC targets",
    "HALLMARK_GLYCOLYSIS": "glycolysis",
    "HALLMARK_OXIDATIVE_PHOSPHORYLATION": "oxidative phosphorylation",
    "GOBP_MITOCHONDRIAL_TRANSLATION": "mitochondrial translation",
    "GOBP_RIBOSOME_BIOGENESIS": "ribosome biogenesis",
    "REACTOME_RRNA_PROCESSING": "rRNA processing",
    "HALLMARK_HYPOXIA": "hypoxia",
    "BUFFA_HYPOXIA_METAGENE": "hypoxia metagene",
    "SEMENZA_HIF1_TARGETS": "HIF1 targets",
    "HALLMARK_ANGIOGENESIS": "angiogenesis",
    "REACTOME_IRON_UPTAKE_AND_TRANSPORT": "iron uptake and transport",
    "HALLMARK_EPITHELIAL_MESENCHYMAL_TRANSITION": "EMT",
    "FRIDMAN_SENESCENCE_UP": "senescence",
    "HALLMARK_MITOTIC_SPINDLE": "mitotic spindle",
    "HALLMARK_G2M_CHECKPOINT": "G2M checkpoint",
    "HALLMARK_E2F_TARGETS": "E2F targets",
    "REACTOME_CITRIC_ACID_CYCLE_TCA_CYCLE": "TCA cycle",
}
by_name = {r["NAME"].strip(): r for r in gsea_rows}
f["themes"] = {}
for k, label in _want.items():
    if k in by_name:
        r = by_name[k]
        f["themes"][label] = {"set": k, "nes": round(r["NES"], 2),
                              "p": r["p"], "q": round(r["q"], 3)}

os.makedirs(FIG, exist_ok=True)
json.dump(f, open(os.path.join(FIG, "facts.json"), "w"), indent=1)
print("facts.json: %d gene sets, %d DE genes (%du/%dd), %d drugs ranked"
      % (f["gsea_n_sets"], f["de_total"], f["de_up"], f["de_down"], len(f["drug_by_name"])))
